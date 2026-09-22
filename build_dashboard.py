"""
build_dashboard.py

Toma el último Excel que produce remaju_scraper_v3.py
(Bloomberg_Remates_ULTIMO.xlsx) y genera docs/index.html:
el dashboard "Remaju Ledger" con los datos ya incrustados.

Se corre automáticamente en el workflow de GitHub Actions,
después de cada barrido del scraper -- así el dashboard publicado
en GitHub Pages queda siempre actualizado con la última corrida,
sin que nadie tenga que tocarlo a mano.

Uso manual (para probar en tu propia compu):
    python build_dashboard.py
"""

import json
import os
import sys
import openpyxl

EXCEL_PATH = "Bloomberg_Remates_ULTIMO.xlsx"
TEMPLATE_PATH = os.path.join("templates", "dashboard_template.html")
OUTPUT_DIR = "docs"
OUTPUT_PATH = os.path.join(OUTPUT_DIR, "index.html")

# Columnas que el dashboard realmente usa. Si en el futuro el
# dashboard necesita un campo nuevo, agrégalo aquí Y en el
# JavaScript de templates/dashboard_template.html (busca las
# referencias a d['Nombre De Columna']).
COLUMNAS_DASHBOARD = [
    "Código de Remate", "Convocatoria", "Tipo Remate", "Fecha del remate",
    "Hora del remate", "Días restantes", "Estado temporal",
    "Número de Expediente", "Distrito Judicial", "Órgano Jurisdisccional",
    "Materia", "Descripción", "N° inscritos", "Tasación", "Moneda Tasación",
    "Precio Base", "Moneda Precio Base", "SUNARP Estado", "SUNARP Cargas",
    "SUNARP Gravámenes", "CEJ Estado Proceso", "Demandante", "Demandado",
    "Gravámenes", "Cargas", "Departamento Inmueble", "Provincia Inmueble",
    "Distrito Inmueble", "Tipo Inmueble 1", "Dirección 1", "Estado REMAJU",
]


def safe(v):
    """Convierte valores de Excel (fechas, etc.) a algo serializable en JSON."""
    if v is None:
        return None
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return v


def cargar_datos_excel(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}

    filas = []
    for r in range(2, ws.max_row + 1):
        fila = {}
        for col in COLUMNAS_DASHBOARD:
            c = idx.get(col)
            fila[col] = safe(ws.cell(row=r, column=c + 1).value) if c is not None else None
        if fila.get("Código de Remate"):
            filas.append(fila)
    return filas


def main():
    if not os.path.exists(EXCEL_PATH):
        print(f"⚠️  No se encontró {EXCEL_PATH} -- el dashboard no se regenera esta corrida.")
        sys.exit(0)  # no es un error fatal: simplemente no hay nada nuevo que publicar

    print(f"📊 Leyendo {EXCEL_PATH}...")
    datos = cargar_datos_excel(EXCEL_PATH)
    print(f"   {len(datos)} remates cargados.")

    if not os.path.exists(TEMPLATE_PATH):
        print(f"❌ No se encontró la plantilla en {TEMPLATE_PATH}.")
        sys.exit(1)

    with open(TEMPLATE_PATH, "r", encoding="utf-8") as f:
        template = f.read()

    data_json = json.dumps(datos, ensure_ascii=False)
    salida = template.replace("__DATA_JSON__", data_json)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        f.write(salida)

    tam_mb = os.path.getsize(OUTPUT_PATH) / 1024 / 1024
    print(f"✅ Dashboard regenerado: {OUTPUT_PATH} ({tam_mb:.2f} MB)")


if __name__ == "__main__":
    main()
